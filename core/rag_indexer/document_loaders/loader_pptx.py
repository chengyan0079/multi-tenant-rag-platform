"""
loader_pptx - PPTX 文档加载器（文本 + 图片 OCR）

Author: 骆昊
"""
from io import BytesIO
from typing import Iterator

import numpy as np
from PIL import Image
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.autoshape import Shape
from tqdm import tqdm

from common.config import config
from common.logger import logger, setup_logging
from core.rag_indexer.document_loaders.ocr import get_ocr_instance


class OCRPPTLoader(BaseLoader):
    """PPT 文档加载器"""

    def __init__(self, filepath: str) -> None:
        self.filepath = filepath

    def lazy_load(self) -> Iterator[Document]:
        text_content = self.pptx2text()
        yield Document(page_content=text_content, metadata={'source': self.filepath})

    def pptx2text(self) -> str:
        presentation = Presentation(self.filepath)
        ocr_obj = get_ocr_instance()
        extracted_chunks: list[str] = []

        def _parse_shape(shape: Shape) -> list[str]:
            chunks = []

            if shape.has_text_frame:
                txt = shape.text.strip()
                if txt:
                    chunks.append(txt)
            elif shape.has_table:
                shape_table = getattr(shape, 'table', None)
                if shape_table:
                    for row in shape_table.rows:
                        for cell in row.cells:
                            if cell.text_frame:
                                cell_txt = cell.text_frame.text.strip()
                                if cell_txt:
                                    chunks.append(cell_txt)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    shape_image = getattr(shape, 'image', None)
                    if shape_image:
                        image_bytes = shape_image.blob
                        image = Image.open(BytesIO(image_bytes)).convert('RGB')
                        result, _ = ocr_obj(np.array(image))
                        if result:
                            ocr_lines = [line[1] for line in result if line[1].strip()]
                            if ocr_lines:
                                chunks.append('\n'.join(ocr_lines))
                except Exception as e:
                    logger.warning(f'图片解析（光学文字识别）错误: {e}', exc_info=True)
            elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                child_shapes = getattr(shape, 'shapes', [])
                for child_shape in child_shapes:
                    chunks.extend(_parse_shape(child_shape))

            return chunks

        pro_bar = tqdm(total=len(presentation.slides), desc='OCRPPTLoader processing')
        for slide in presentation.slides:
            sorted_shapes = sorted(slide.shapes, key=lambda x: (x.top or 0, x.left or 0))
            for curr_shape in sorted_shapes:
                shape_text_list = _parse_shape(curr_shape)
                extracted_chunks.extend(shape_text_list)
            pro_bar.update(1)

        pro_bar.close()
        return '\n'.join(extracted_chunks)


if __name__ == '__main__':
    setup_logging()
    ppt_loader = OCRPPTLoader(filepath=str(config.FILES_DIR / '英语面试指南.pptx'))
    docs = ppt_loader.load()
    logger.info(docs)
    logger.info(len(docs[0].page_content))
