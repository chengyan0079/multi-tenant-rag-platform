"""
loader_docx - DOCX 文档加载器（文本 + 嵌入图片 OCR）

Author: 骆昊
"""
from io import BytesIO
from typing import Iterator

import docx
import numpy as np
from PIL import Image
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.parts.image import ImagePart
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document
from tqdm import tqdm

from common.config import config
from common.logger import logger, setup_logging
from core.rag_indexer.document_loaders.ocr import get_ocr_instance


def _iter_block_items(parent: DocxDocument | _Cell) -> Iterator[Paragraph | Table]:
    if isinstance(parent, DocxDocument):
        parent_elem = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elem = parent._tc
    else:
        raise TypeError(f'不支持的父元素解析类型: {type(parent)}')

    for child in parent_elem.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


class OCRDocxLoader(BaseLoader):
    """支持文本提取与嵌入图片 OCR 识别的 Word 加载器"""

    def __init__(self, filepath: str) -> None:
        self.filepath = filepath

    def lazy_load(self) -> Iterator[Document]:
        text_content = self.docx2text()
        yield Document(page_content=text_content, metadata={'source': self.filepath})

    def docx2text(self) -> str:
        ocr_obj = get_ocr_instance()
        doc = docx.Document(self.filepath)
        extracted_chunks: list[str] = []

        total_blocks = len(doc.paragraphs) + len(doc.tables)
        pro_bar = tqdm(total=total_blocks, desc='OCRDOCLoader processing')

        for block in _iter_block_items(doc):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if text:
                    extracted_chunks.append(text)

                pic_elements = block._element.xpath('.//pic:pic')
                for pic in pic_elements:
                    for embed_id in pic.xpath('.//a:blip/@r:embed'):
                        part = doc.part.related_parts.get(embed_id)
                        if isinstance(part, ImagePart):
                            try:
                                with BytesIO(part.blob) as image_stream:
                                    image = Image.open(image_stream).convert('RGB')
                                    result, _ = ocr_obj(np.array(image))
                                    if result:
                                        ocr_lines = [line[1] for line in result if line[1].strip()]
                                        if ocr_lines:
                                            extracted_chunks.append("\n".join(ocr_lines))
                            except Exception as e:
                                logger.warning(f'提取 Word 文档中图像文字发生错误: {e}', exc_info=True)

            elif isinstance(block, Table):
                for row in block.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            cell_text = paragraph.text.strip()
                            if cell_text:
                                extracted_chunks.append(cell_text)

            pro_bar.update(1)

        pro_bar.close()
        return '\n'.join(extracted_chunks)


if __name__ == '__main__':
    setup_logging()
    docx_loader = OCRDocxLoader(filepath=str(config.FILES_DIR / '人工智能就业课课程大纲.docx'))
    docs = docx_loader.load()
    logger.info(docs)
    logger.info(len(docs[0].page_content))
